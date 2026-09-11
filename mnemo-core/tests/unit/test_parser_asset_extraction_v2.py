"""Phase 8.5.2 parser asset-discovery and provenance contracts."""

from __future__ import annotations

import base64
import hashlib
import io
import zipfile
from dataclasses import replace
from datetime import UTC, datetime

import docx
import fitz  # type: ignore[import-untyped]
import pytest
from mnemo.interfaces.errors import ContractValidationError
from mnemo.interfaces.parser_models import (
    AssetExtractionOmission,
    AssetExtractionOutcome,
    AssetOmissionReason,
    ParseResult,
    ParseResultV2,
    RawImageBlock,
    RawTextBlock,
    TransientAsset,
    TransientAssetOccurrence,
)
from mnemo.interfaces.types import FileMetadata
from mnemo.models import (
    AssetContainerKind,
    AssetExtractionProvenance,
    AssetLocator,
    AssetLocatorKind,
    DocType,
    DocumentMetadata,
    FrozenMetadata,
)
from mnemo.parsers.asset_extraction import (
    AssetExtractionLimits,
    asset_limit_reason,
    bounded_asset_result,
    extraction_from_parse_result,
    image_signature_media_type,
    resolve_archive_target,
    safe_xml_from_bytes,
    validate_image_payload,
    validated_zip,
)
from mnemo.parsers.docx import DOCXParser
from mnemo.parsers.html import HTMLParser
from mnemo.parsers.image import StandaloneImageParser
from mnemo.parsers.markdown import MarkdownParser
from mnemo.parsers.pdf import PDFParser
from mnemo.parsers.pptx import PPTXParser
from mnemo.parsers.xlsx import XLSXParser
from openpyxl import Workbook  # type: ignore[import-untyped]

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _metadata(data: bytes, mime_type: str) -> FileMetadata:
    return FileMetadata(
        content_hash=hashlib.sha256(data).hexdigest(),
        size_bytes=len(data),
        mime_type=mime_type,
        modified_at=datetime(2026, 8, 24, tzinfo=UTC),
        metadata=FrozenMetadata(),
    )


def _pptx(*, missing_relationship: bool = False, external_relationship: bool = False) -> bytes:
    slide = """<?xml version="1.0" encoding="UTF-8"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
 xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
 <p:cSld><p:spTree>
  <p:pic><p:nvPicPr><p:cNvPr id="2" name="Picture 1" descr="diagram"/></p:nvPicPr>
   <p:blipFill><a:blip r:embed="rId2"/></p:blipFill>
   <p:spPr><a:xfrm><a:off x="10" y="20"/><a:ext cx="30" cy="40"/></a:xfrm></p:spPr>
  </p:pic>
  <p:pic><p:nvPicPr><p:cNvPr id="3" name="Picture 2"/></p:nvPicPr>
   <p:blipFill><a:blip r:embed="rId2"/></p:blipFill><p:spPr/>
  </p:pic>
 </p:spTree></p:cSld>
</p:sld>"""
    relationship = ""
    if not missing_relationship:
        target = (
            "https://example.invalid/image.png" if external_relationship else "../media/image1.png"
        )
        mode = ' TargetMode="External"' if external_relationship else ""
        relationship = (
            f'<Relationship Id="rId2" Target="{target}" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            f'relationships/image"{mode}/>'
        )
    rels = f"""<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{relationship}</Relationships>"""
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("ppt/presentation.xml", "<presentation/>")
        archive.writestr("ppt/slides/slide1.xml", slide)
        archive.writestr("ppt/slides/_rels/slide1.xml.rels", rels)
        archive.writestr("ppt/media/image1.png", PNG)
    return output.getvalue()


def _xlsx_with_image() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Evidence"
    sheet.append(("label", "value"))
    sheet.append(("answer", 42))
    base = io.BytesIO()
    workbook.save(base)
    workbook.close()

    input_zip = zipfile.ZipFile(io.BytesIO(base.getvalue()))
    sheet_xml = input_zip.read("xl/worksheets/sheet1.xml").decode("utf-8")
    sheet_xml = sheet_xml.replace(
        "</worksheet>",
        '<drawing xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
        'relationships" r:id="rId9"/></worksheet>',
    )
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for info in input_zip.infolist():
            if info.filename != "xl/worksheets/sheet1.xml":
                archive.writestr(info, input_zip.read(info.filename))
        archive.writestr("xl/worksheets/sheet1.xml", sheet_xml)
        archive.writestr(
            "xl/worksheets/_rels/sheet1.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
            'relationships"><Relationship Id="rId9" Target="../drawings/drawing1.xml"/>'
            "</Relationships>",
        )
        archive.writestr(
            "xl/drawings/drawing1.xml",
            '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/'
            'spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/'
            '2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/'
            '2006/relationships"><xdr:twoCellAnchor><xdr:from><xdr:col>1</xdr:col>'
            "<xdr:row>2</xdr:row></xdr:from><xdr:to><xdr:col>3</xdr:col><xdr:row>5"
            '</xdr:row></xdr:to><xdr:pic><xdr:blipFill><a:blip r:embed="rId1"/>'
            "</xdr:blipFill></xdr:pic></xdr:twoCellAnchor></xdr:wsDr>",
        )
        archive.writestr(
            "xl/drawings/_rels/drawing1.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
            'relationships"><Relationship Id="rId1" Target="../media/image1.png"/>'
            "</Relationships>",
        )
        archive.writestr("xl/media/image1.png", PNG)
    input_zip.close()
    return output.getvalue()


def _rewrite_zip(
    data: bytes, *, replacements: dict[str, bytes], omitted: set[str] = frozenset()
) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for info in source.infolist():
            if info.filename not in omitted:
                archive.writestr(info, replacements.get(info.filename, source.read(info.filename)))
    source.close()
    return output.getvalue()


def test_standalone_image_preserves_exact_bytes_and_provenance() -> None:
    result = StandaloneImageParser().parse_with_assets(
        PNG, "diagram.png", _metadata(PNG, "image/png")
    )
    assert result.outcome is AssetExtractionOutcome.COMPLETE
    assert result.parse_result.extracted_assets[0].raw_bytes == PNG
    assert result.parse_result.extracted_assets[0].mime_type == "image/png"
    assert result.asset_occurrences[0].locator.kind is AssetLocatorKind.STANDALONE
    assert result.asset_occurrences[0].extraction_provenance.parser_id == "mnemo.image"


def test_standalone_image_v1_capabilities_and_empty_validation() -> None:
    parser = StandaloneImageParser()
    assert ".png" in parser.supported_formats
    assert parser.capabilities().supports_images is True
    assert parser.parse(PNG, "diagram.png", _metadata(PNG, "image/png")).extracted_assets
    with pytest.raises(ContractValidationError, match="empty image"):
        parser.parse_with_assets(b"", "empty.png", _metadata(b"", "image/png"))


def test_standalone_image_rejects_mime_spoof_and_active_svg() -> None:
    with pytest.raises(ContractValidationError, match="does not match"):
        StandaloneImageParser().parse_with_assets(PNG, "fake.jpg", _metadata(PNG, "image/jpeg"))
    svg = b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>'
    with pytest.raises(ContractValidationError, match="active content"):
        StandaloneImageParser().parse_with_assets(
            svg, "active.svg", _metadata(svg, "image/svg+xml")
        )


def test_pdf_extracts_repeated_bytes_as_distinct_page_occurrences() -> None:
    document = fitz.open()
    for _ in range(2):
        page = document.new_page()
        page.insert_image(fitz.Rect(10, 20, 30, 40), stream=PNG)
    data = document.write()
    document.close()
    result = PDFParser().parse_with_assets(data, "paper.pdf", _metadata(data, "application/pdf"))
    assert result.outcome is AssetExtractionOutcome.COMPLETE
    assert [item.locator.page_number for item in result.asset_occurrences] == [1, 2]
    assert len({item.parser_local_id for item in result.asset_occurrences}) == 2
    extracted = {asset.raw_bytes for asset in result.parse_result.extracted_assets}
    assert len(extracted) == 1
    assert next(iter(extracted)).startswith(b"\x89PNG\r\n\x1a\n")


def test_pdf_without_images_reports_no_assets_and_image_only_has_no_ocr_text() -> None:
    document = fitz.open()
    document.new_page()
    data = document.write()
    document.close()
    result = PDFParser().parse_with_assets(data, "empty.pdf", _metadata(data, "application/pdf"))
    assert result.outcome is AssetExtractionOutcome.NO_ASSETS
    assert result.parse_result.blocks == ()


def test_docx_preserves_repeated_relationship_occurrences_and_v1_projection() -> None:
    document = docx.Document()
    paragraph = document.add_paragraph("Evidence ")
    paragraph.add_run().add_picture(io.BytesIO(PNG))
    paragraph.add_run().add_picture(io.BytesIO(PNG))
    output = io.BytesIO()
    document.save(output)
    data = output.getvalue()
    parser = DOCXParser()
    result = parser.parse_with_assets(
        data,
        "report.docx",
        _metadata(data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    )
    assert result.outcome is AssetExtractionOutcome.COMPLETE
    assert len(result.asset_occurrences) == 2
    assert all(
        item.locator.kind is AssetLocatorKind.DOCX_POSITION for item in result.asset_occurrences
    )
    assert (
        len(
            parser.parse(
                data,
                "report.docx",
                _metadata(
                    data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                ),
            ).extracted_assets
        )
        == 1
    )


def test_docx_corrupt_optional_image_does_not_discard_text() -> None:
    document = docx.Document()
    document.add_paragraph("Text must survive")
    document.add_picture(io.BytesIO(PNG))
    source = io.BytesIO()
    document.save(source)
    original = zipfile.ZipFile(io.BytesIO(source.getvalue()))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for info in original.infolist():
            payload = original.read(info.filename)
            if info.filename.startswith("word/media/"):
                payload = b"corrupt image"
            archive.writestr(info, payload)
    original.close()
    data = output.getvalue()
    result = DOCXParser().parse_with_assets(
        data,
        "report.docx",
        _metadata(data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    )
    assert result.outcome is AssetExtractionOutcome.REJECTED
    assert result.omissions[0].reason is AssetOmissionReason.CORRUPT_ASSET
    assert any(
        isinstance(block, RawTextBlock) and block.text == "Text must survive"
        for block in result.parse_result.blocks
    )


def test_pptx_extracts_repeated_slide_occurrences_without_changing_v1_text() -> None:
    data = _pptx()
    parser = PPTXParser()
    metadata = _metadata(
        data, "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    result = parser.parse_with_assets(data, "slides.pptx", metadata)
    assert result.outcome is AssetExtractionOutcome.COMPLETE
    assert len(result.asset_occurrences) == 2
    assert result.asset_occurrences[0].locator.slide_number == 1
    assert result.asset_occurrences[0].locator.geometry == (10.0, 20.0, 40.0, 60.0)
    assert result.asset_occurrences[0].authored_alt_text == "diagram"
    assert parser.parse(data, "slides.pptx", metadata).extracted_assets == ()


def test_pptx_retains_non_picture_image_relationships_and_jpeg_xr() -> None:
    data = _pptx()
    source = zipfile.ZipFile(io.BytesIO(data))
    rel_path = "ppt/slides/_rels/slide1.xml.rels"
    relationships = (
        source.read(rel_path)
        .decode("utf-8")
        .replace(
            "</Relationships>",
            '<Relationship Id="rId3" Target="../media/image2.wdp" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hdphoto"/>'
            "</Relationships>",
        )
    )
    source.close()
    jpeg_xr = b"II\xbc\x01" + b"\0" * 32
    data = _rewrite_zip(data, replacements={rel_path: relationships.encode()})
    output = io.BytesIO()
    rewritten = zipfile.ZipFile(io.BytesIO(data))
    with zipfile.ZipFile(output, "w") as archive:
        for info in rewritten.infolist():
            archive.writestr(info, rewritten.read(info.filename))
        archive.writestr("ppt/media/image2.wdp", jpeg_xr)
    rewritten.close()

    result = PPTXParser().parse_with_assets(
        output.getvalue(),
        "slides.pptx",
        _metadata(output.getvalue(), "application/vnd.ms-powerpoint"),
    )

    assert len(result.asset_occurrences) == 3
    assert result.parse_result.extracted_assets[-1].raw_bytes == jpeg_xr
    assert result.parse_result.extracted_assets[-1].mime_type == "image/vnd.ms-photo"
    fallback = result.parse_result.blocks[-1]
    assert isinstance(fallback, RawImageBlock)
    assert fallback.metadata["parser.asset.package_fallback"] is True


def test_pptx_missing_relationship_is_auditable_not_fatal() -> None:
    data = _pptx(missing_relationship=True)
    result = PPTXParser().parse_with_assets(
        data,
        "slides.pptx",
        _metadata(
            data, "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        ),
    )
    assert result.outcome is AssetExtractionOutcome.REJECTED
    assert {item.reason for item in result.omissions} == {
        AssetOmissionReason.MALFORMED_RELATIONSHIP
    }


def test_pptx_external_relationship_is_never_fetched() -> None:
    data = _pptx(external_relationship=True)
    result = PPTXParser().parse_with_assets(
        data,
        "slides.pptx",
        _metadata(
            data, "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        ),
    )
    assert result.outcome is AssetExtractionOutcome.REJECTED
    assert {item.reason for item in result.omissions} == {AssetOmissionReason.EXTERNAL_REFERENCE}


def test_pptx_asset_count_limit_is_applied_before_second_archive_read() -> None:
    data = _pptx()
    result = PPTXParser(AssetExtractionLimits(max_assets=1)).parse_with_assets(
        data,
        "slides.pptx",
        _metadata(
            data, "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        ),
    )
    assert result.outcome is AssetExtractionOutcome.PARTIAL
    assert len(result.asset_occurrences) == 1
    assert result.omissions[0].reason is AssetOmissionReason.LIMIT_EXCEEDED


def test_xlsx_extracts_drawing_with_sheet_and_anchor_provenance() -> None:
    data = _xlsx_with_image()
    parser = XLSXParser()
    metadata = _metadata(data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    result = parser.parse_with_assets(data, "workbook.xlsx", metadata)
    assert result.outcome is AssetExtractionOutcome.COMPLETE
    locator = result.asset_occurrences[0].locator
    assert locator.kind is AssetLocatorKind.XLSX_CELL
    assert locator.sheet_name == "Evidence"
    assert locator.cell_reference == "B3:D6"
    assert parser.parse(data, "workbook.xlsx", metadata).extracted_assets == ()


@pytest.mark.parametrize(
    ("replacements", "omitted", "reason"),
    [
        (
            {
                "xl/worksheets/_rels/sheet1.xml.rels": (
                    b'<Relationships xmlns="http://schemas.openxmlformats.org/package/'
                    b'2006/relationships"/>'
                )
            },
            frozenset(),
            AssetOmissionReason.MALFORMED_RELATIONSHIP,
        ),
        (
            {},
            frozenset({"xl/drawings/drawing1.xml"}),
            AssetOmissionReason.MALFORMED_RELATIONSHIP,
        ),
        (
            {
                "xl/drawings/_rels/drawing1.xml.rels": (
                    b'<Relationships xmlns="http://schemas.openxmlformats.org/package/'
                    b'2006/relationships"><Relationship Id="rId1" '
                    b'Target="https://example.invalid/image.png" TargetMode="External"/>'
                    b"</Relationships>"
                )
            },
            frozenset(),
            AssetOmissionReason.EXTERNAL_REFERENCE,
        ),
    ],
)
def test_xlsx_relationship_failures_are_typed_and_text_survives(
    replacements: dict[str, bytes], omitted: frozenset[str], reason: AssetOmissionReason
) -> None:
    data = _rewrite_zip(_xlsx_with_image(), replacements=replacements, omitted=set(omitted))
    result = XLSXParser().parse_with_assets(
        data,
        "workbook.xlsx",
        _metadata(data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    )
    assert result.outcome is AssetExtractionOutcome.REJECTED
    assert reason in {item.reason for item in result.omissions}
    assert any(
        isinstance(block, RawTextBlock) or block.__class__.__name__ == "RawTableBlock"
        for block in result.parse_result.blocks
    )


def test_data_uri_html_and_markdown_emit_dom_occurrences_without_network_fetch() -> None:
    encoded = base64.b64encode(PNG).decode("ascii")
    html = (
        f'<html><body><p>Text</p><img alt="plot" src="data:image/png;base64,{encoded}">'
        "</body></html>"
    ).encode()
    html_result = HTMLParser().parse_with_assets(html, "page.html", _metadata(html, "text/html"))
    assert html_result.outcome is AssetExtractionOutcome.COMPLETE
    assert html_result.asset_occurrences[0].locator.kind is AssetLocatorKind.DOM_POSITION
    markdown = f"Text ![plot](data:image/png;base64,{encoded})".encode()
    markdown_result = MarkdownParser().parse_with_assets(
        markdown, "page.md", _metadata(markdown, "text/markdown")
    )
    assert markdown_result.outcome is AssetExtractionOutcome.COMPLETE
    assert markdown_result.asset_occurrences[0].locator.kind is AssetLocatorKind.DOM_POSITION


def test_archive_and_asset_limits_fail_closed_without_path_traversal() -> None:
    malicious = io.BytesIO()
    with zipfile.ZipFile(malicious, "w") as archive:
        archive.writestr("../escape.png", PNG)
    with pytest.raises(ContractValidationError, match="path is unsafe"):
        validated_zip(malicious.getvalue(), AssetExtractionLimits())

    document = docx.Document()
    document.add_picture(io.BytesIO(PNG))
    output = io.BytesIO()
    document.save(output)
    data = output.getvalue()
    result = DOCXParser(
        AssetExtractionLimits(max_asset_bytes=1, max_total_asset_bytes=1)
    ).parse_with_assets(
        data,
        "bounded.docx",
        _metadata(data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    )
    assert result.outcome is AssetExtractionOutcome.REJECTED
    assert result.omissions[0].reason is AssetOmissionReason.LIMIT_EXCEEDED
    assert not any(isinstance(block, RawImageBlock) for block in result.parse_result.blocks)


@pytest.mark.parametrize(
    ("payload", "expected"),
    [
        (b"\xff\xd8\xffjpeg", "image/jpeg"),
        (b"GIF89aimage", "image/gif"),
        (b"II*\x00tiff", "image/tiff"),
        (b"BMbitmap", "image/bmp"),
        (b"RIFFxxxxWEBPdata", "image/webp"),
        (b"  <svg xmlns='http://www.w3.org/2000/svg'/>", "image/svg+xml"),
        (b"not-an-image", None),
    ],
)
def test_image_signature_detection_is_format_generic(payload: bytes, expected: str | None) -> None:
    assert image_signature_media_type(payload) == expected


def test_image_payload_bounds_and_declared_aliases() -> None:
    limits = AssetExtractionLimits(max_asset_bytes=len(PNG), max_total_asset_bytes=len(PNG))
    assert validate_image_payload(PNG, "image/x-png; charset=binary", limits) == "image/png"
    with pytest.raises(ContractValidationError, match="empty"):
        validate_image_payload(b"", "image/png", limits)
    with pytest.raises(ContractValidationError, match="byte limit"):
        validate_image_payload(PNG + b"x", "image/png", limits)
    with pytest.raises(ContractValidationError, match="unsupported or malformed"):
        validate_image_payload(b"unknown", "application/octet-stream", limits)


def test_limits_and_xml_relationship_security_validation() -> None:
    with pytest.raises(ValueError, match="positive integer"):
        AssetExtractionLimits(max_assets=0)
    with pytest.raises(ValueError, match="cannot be smaller"):
        AssetExtractionLimits(max_asset_bytes=2, max_total_asset_bytes=1)
    with pytest.raises(ValueError, match="greater than one"):
        AssetExtractionLimits(max_compression_ratio=1)
    with pytest.raises(ContractValidationError, match="valid ZIP"):
        validated_zip(b"broken", AssetExtractionLimits())
    with pytest.raises(ContractValidationError, match="invalid"):
        resolve_archive_target("word/document.xml", "..\\media\\image.png")
    with pytest.raises(ContractValidationError, match="external"):
        resolve_archive_target("word/document.xml", "https://example.invalid/image.png")
    assert resolve_archive_target("xl/workbook.xml", "/xl/media/image.png") == (
        "xl/media/image.png"
    )
    with pytest.raises(ContractValidationError, match="entity"):
        safe_xml_from_bytes(b"<!DOCTYPE x [<!ENTITY y 'z'>]><x/>")
    with pytest.raises(ContractValidationError, match="malformed XML"):
        safe_xml_from_bytes(b"<x>")
    assert (
        asset_limit_reason(
            asset_count=0,
            total_bytes=5,
            next_bytes=6,
            limits=AssetExtractionLimits(max_asset_bytes=10, max_total_asset_bytes=10),
        )
        is AssetOmissionReason.LIMIT_EXCEEDED
    )
    assert (
        asset_limit_reason(
            asset_count=0,
            total_bytes=0,
            next_bytes=1,
            limits=AssetExtractionLimits(),
        )
        is None
    )


def test_archive_count_expansion_and_compression_bounds() -> None:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("one", "1")
        archive.writestr("two", "2")
    with pytest.raises(ContractValidationError, match="entry-count"):
        validated_zip(output.getvalue(), AssetExtractionLimits(max_archive_entries=1))
    with pytest.raises(ContractValidationError, match="expansion"):
        validated_zip(
            output.getvalue(),
            AssetExtractionLimits(max_archive_uncompressed_bytes=1),
        )

    compressed = io.BytesIO()
    with zipfile.ZipFile(compressed, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("repeated", b"0" * 4096)
    with pytest.raises(ContractValidationError, match="compression-ratio"):
        validated_zip(
            compressed.getvalue(),
            AssetExtractionLimits(max_compression_ratio=2),
        )


def test_total_asset_and_count_bounds_are_auditable() -> None:
    parsed = _raw_result(("one", PNG, "image/png"), ("two", PNG, "image/png"))
    count_limited = bounded_asset_result(
        parsed,
        parser_id="synthetic",
        container_kind=AssetContainerKind.OTHER,
        limits=AssetExtractionLimits(max_assets=1),
    )
    assert count_limited.outcome is AssetExtractionOutcome.PARTIAL
    assert count_limited.omissions[0].reason is AssetOmissionReason.LIMIT_EXCEEDED
    byte_limited = bounded_asset_result(
        parsed,
        parser_id="synthetic",
        container_kind=AssetContainerKind.OTHER,
        limits=AssetExtractionLimits(
            max_asset_bytes=len(PNG),
            max_total_asset_bytes=len(PNG) + 1,
        ),
    )
    assert byte_limited.outcome is AssetExtractionOutcome.PARTIAL
    assert byte_limited.omissions[0].reason is AssetOmissionReason.LIMIT_EXCEEDED


def _raw_result(*assets: tuple[str, bytes, str]) -> ParseResult:
    blocks = [RawTextBlock(ordinal=0, text="canonical text")]
    transients = []
    for index, (local_id, payload, mime_type) in enumerate(assets, start=1):
        blocks.append(
            RawImageBlock(
                ordinal=index,
                parser_local_id=local_id,
                metadata=FrozenMetadata(
                    {
                        "parser.asset.relationship_id": f"rId{index}",
                        "parser.asset.dom_path": f"body/img[{index}]",
                    }
                ),
            )
        )
        transients.append(
            TransientAsset(
                parser_local_id=local_id,
                raw_bytes=payload,
                mime_type=mime_type,
            )
        )
    return ParseResult(
        blocks=tuple(blocks),
        extracted_assets=tuple(transients),
        metadata=DocumentMetadata(content_hash="a" * 64),
        language="und",
        doc_type=DocType.GENERIC,
    )


def test_bounded_result_is_partial_and_never_changes_canonical_text() -> None:
    parsed = _raw_result(("good", PNG, "application/octet-stream"), ("bad", b"bad", "image/png"))
    result = bounded_asset_result(
        parsed,
        parser_id="synthetic",
        container_kind=AssetContainerKind.HTML,
        limits=AssetExtractionLimits(),
    )
    assert result.outcome is AssetExtractionOutcome.PARTIAL
    assert result.parse_result.blocks[0] == parsed.blocks[0]
    assert result.parse_result.blocks[0].text == "canonical text"
    assert result.parse_result.extracted_assets[0].mime_type == "image/png"
    assert result.asset_occurrences[0].locator.kind is AssetLocatorKind.DOM_POSITION
    assert result.omissions[0].reason is AssetOmissionReason.CORRUPT_ASSET
    assert result.omissions[0].relationship_id == "rId2"


@pytest.mark.parametrize(
    ("container", "metadata", "expected"),
    [
        (
            AssetContainerKind.DOCX,
            FrozenMetadata({"parser.asset.inline_position": "invalid"}),
            AssetLocatorKind.DOCX_POSITION,
        ),
        (
            AssetContainerKind.XLSX,
            FrozenMetadata({"parser.asset.sheet_name": "Sheet1"}),
            AssetLocatorKind.XLSX_CELL,
        ),
        (AssetContainerKind.OTHER, FrozenMetadata(), AssetLocatorKind.DOCUMENT_POSITION),
    ],
)
def test_generic_locator_fallbacks_are_typed(
    container: AssetContainerKind,
    metadata: FrozenMetadata,
    expected: AssetLocatorKind,
) -> None:
    parsed = ParseResult(
        blocks=(RawImageBlock(ordinal=0, parser_local_id="image", metadata=metadata),),
        extracted_assets=(
            TransientAsset(parser_local_id="image", raw_bytes=PNG, mime_type="image/png"),
        ),
        metadata=DocumentMetadata(content_hash="b" * 64),
        language="und",
        doc_type=DocType.GENERIC,
    )
    assert (
        extraction_from_parse_result(parsed, parser_id="synthetic", container_kind=container)
        .asset_occurrences[0]
        .locator.kind
        is expected
    )


def test_parse_result_v2_rejects_inconsistent_state() -> None:
    parsed = _raw_result(("asset", PNG, "image/png"))
    omission = AssetExtractionOmission(reason=AssetOmissionReason.CORRUPT_ASSET, ordinal=0)
    with pytest.raises(ValueError, match="COMPLETE"):
        ParseResultV2(
            parse_result=ParseResult(
                blocks=(),
                extracted_assets=(),
                metadata=DocumentMetadata(content_hash="c" * 64),
                language="und",
                doc_type=DocType.GENERIC,
            ),
            asset_occurrences=(),
            omissions=(),
            outcome=AssetExtractionOutcome.COMPLETE,
        )
    with pytest.raises(ValueError, match="NO_ASSETS"):
        ParseResultV2(
            parse_result=parsed,
            asset_occurrences=(),
            omissions=(omission,),
            outcome=AssetExtractionOutcome.NO_ASSETS,
        )
    with pytest.raises(ValueError, match="PARTIAL"):
        ParseResultV2(
            parse_result=parsed,
            asset_occurrences=(),
            omissions=(omission,),
            outcome=AssetExtractionOutcome.PARTIAL,
        )
    with pytest.raises(ValueError, match="REJECTED"):
        ParseResultV2(
            parse_result=parsed,
            asset_occurrences=(),
            omissions=(),
            outcome=AssetExtractionOutcome.REJECTED,
        )
    with pytest.raises(TypeError, match="parse_result"):
        ParseResultV2(  # type: ignore[arg-type]
            parse_result="bad",
            asset_occurrences=(),
            omissions=(),
            outcome=AssetExtractionOutcome.NO_ASSETS,
        )
    with pytest.raises(TypeError, match="asset_occurrences"):
        ParseResultV2(  # type: ignore[arg-type]
            parse_result=parsed,
            asset_occurrences=[],
            omissions=(),
            outcome=AssetExtractionOutcome.NO_ASSETS,
        )
    with pytest.raises(TypeError, match="omissions"):
        ParseResultV2(  # type: ignore[arg-type]
            parse_result=parsed,
            asset_occurrences=(),
            omissions=[],
            outcome=AssetExtractionOutcome.NO_ASSETS,
        )
    with pytest.raises(TypeError, match="outcome"):
        ParseResultV2(  # type: ignore[arg-type]
            parse_result=parsed,
            asset_occurrences=(),
            omissions=(),
            outcome="bad",
        )
    valid = extraction_from_parse_result(
        parsed, parser_id="synthetic", container_kind=AssetContainerKind.STANDALONE
    ).asset_occurrences[0]
    with pytest.raises(ValueError, match="unknown transient asset"):
        ParseResultV2(
            parse_result=parsed,
            asset_occurrences=(replace(valid, parser_local_id="missing"),),
            omissions=(),
            outcome=AssetExtractionOutcome.COMPLETE,
        )
    with pytest.raises(TypeError, match="TransientAssetOccurrence"):
        ParseResultV2(  # type: ignore[arg-type]
            parse_result=parsed,
            asset_occurrences=("bad",),
            omissions=(),
            outcome=AssetExtractionOutcome.COMPLETE,
        )
    with pytest.raises(TypeError, match="AssetExtractionOmission"):
        ParseResultV2(  # type: ignore[arg-type]
            parse_result=parsed,
            asset_occurrences=(),
            omissions=("bad",),
            outcome=AssetExtractionOutcome.REJECTED,
        )


def test_occurrence_and_omission_transport_validation() -> None:
    locator = AssetLocator(kind=AssetLocatorKind.STANDALONE, ordinal=0)
    provenance = AssetExtractionProvenance(parser_id="synthetic", parser_version="v2")
    with pytest.raises(ValueError, match="parser_local_id"):
        TransientAssetOccurrence(
            parser_local_id="",
            locator=locator,
            authored_alt_text=None,
            extraction_provenance=provenance,
        )
    with pytest.raises(TypeError, match="locator"):
        TransientAssetOccurrence(  # type: ignore[arg-type]
            parser_local_id="image",
            locator="bad",
            authored_alt_text=None,
            extraction_provenance=provenance,
        )
    with pytest.raises(ValueError, match="authored_alt_text"):
        TransientAssetOccurrence(
            parser_local_id="image",
            locator=locator,
            authored_alt_text="",
            extraction_provenance=provenance,
        )
    with pytest.raises(TypeError, match="extraction_provenance"):
        TransientAssetOccurrence(  # type: ignore[arg-type]
            parser_local_id="image",
            locator=locator,
            authored_alt_text=None,
            extraction_provenance="bad",
        )
    with pytest.raises(TypeError, match="reason"):
        AssetExtractionOmission(reason="bad", ordinal=0)  # type: ignore[arg-type]
    with pytest.raises(ValueError, match="ordinal"):
        AssetExtractionOmission(reason=AssetOmissionReason.CORRUPT_ASSET, ordinal=-1)
    with pytest.raises(TypeError, match="locator"):
        AssetExtractionOmission(  # type: ignore[arg-type]
            reason=AssetOmissionReason.CORRUPT_ASSET, ordinal=0, locator="bad"
        )
    with pytest.raises(ValueError, match="relationship_id"):
        AssetExtractionOmission(
            reason=AssetOmissionReason.CORRUPT_ASSET, ordinal=0, relationship_id=""
        )
